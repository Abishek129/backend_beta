"""Build the EMS tree from a spreadsheet — the upload behind ``POST /api/import/``.

The column lists below ARE the format: ``build_template()`` writes them out and
``run_import()`` reads them back. Two sheets:

    Meters  (required)  one row per meter per place it appears
    Panels  (optional)  one row per panel — panel nameplate and sort order

Add/update only: an upload never deletes, retires or unlinks anything. The
whole upload is one transaction; a preview runs the exact same code and rolls
back, so what it reports is exactly what a commit would do.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from django.db import transaction

from core import logger_devices
from core.logger_client import LoggerApiError
from core.models import (Group, LoggerSchema, Membership, MFM, MfmConfig,
                         NameplateBase, NavCategory, PanelConfig)
from core.nav import category_label, panel_axes, resolved_category
from core.slugs import free_slug, slug_owner

MAX_ROWS = 5000
METERS, PANELS = 'Meters', 'Panels'


# --------------------------------------------------------------------------
# the format
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Column:
    key: str                  # model field for nameplate columns, else an internal key
    label: str                # the header written into the template
    kind: str = 'text'        # text | category | role | number | int
    required: bool = False
    note: str = ''


METER_COLUMNS = (
    Column('name', 'MFM name', required=True,
           note='The meter. The same MFM name + Table name on two rows is the same meter '
                '(e.g. one panel\'s outgoer that is another panel\'s incomer).'),
    Column('category', 'Nav category', 'category', required=True,
           note='One of the EMS nav categories (Lists sheet). With a Distribution '
                'panel it is where the PANEL shows; without one, where this MFM shows.'),
    Column('panel', 'Distribution panel',
           note='Blank = this MFM shows directly under the Nav category. A name = the panel '
                '(created the first time it appears).'),
    Column('role', 'Role', 'role',
           note='Incomings, Outgoings or Spare. Required with a Distribution panel, blank without.'),
    Column('table_name', 'Table name',
           note='The LoggerFast table this meter reads; links its LoggerFast device and data '
                'source. Blank = no live data yet.'),
    Column('logger_schema', 'Schema',
           note='Versioned LoggerFast field layout, for example sch_EMMS7_96. Blank leaves an '
                'existing assignment unchanged.'),
    Column('nominal_voltage_v', 'Nominal voltage (V)', 'number',
           note='Line-to-line, in volts: 415, 11000, 33000.'),
    Column('rated_kva', 'Rated kVA', 'number'),
    Column('rated_kw', 'Rated kW', 'number'),
    Column('rated_current_a', 'Rated current (A)', 'number'),
    Column('frequency_hz', 'Frequency (Hz)', 'number'),
    Column('contracted_kw', 'Contracted kW', 'number'),
)

PANEL_COLUMNS = (
    Column('name', 'Panel name', required=True,
           note='Must match the Distribution panel name used on the Meters sheet.'),
    Column('category', 'Nav category', 'category',
           note='Optional when the Meters sheet already places the panel (it must then match). '
                'Required for a panel with no meters yet.'),
    Column('sort_order', 'Sort order', 'int',
           note='Position inside its category, lower first. Default 0, then by name.'),
    Column('nominal_voltage_v', 'Nominal voltage (V)', 'number',
           note='Bus voltage, line-to-line, in volts.'),
    Column('rated_kva', 'Rated kVA', 'number'),
    Column('rated_kw', 'Rated kW', 'number'),
    Column('rated_current_a', 'Rated current (A)', 'number'),
    Column('fault_level_ka', 'Fault level (kA)', 'number'),
    Column('frequency_hz', 'Frequency (Hz)', 'number'),
    Column('contracted_kw', 'Contracted kW', 'number'),
    Column('critical_load_kw', 'Critical load (kW)', 'number'),
)

# Any other nameplate field may be added as an extra column headed by its field
# name (e.g. "thd_v_limit_pct"); the template just does not list them all.
NAMEPLATE_FIELDS = frozenset(NameplateBase.CONFIG_FIELDS)


def _norm(value):
    """Header / enum matching: case, spaces and punctuation do not matter."""
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).strip()


def _name_key(value):
    """Identity of a meter or panel NAME: case and repeated spaces do not matter."""
    return ' '.join(str(value or '').split()).casefold()


def _short(label):
    """'HT Panels' -> 'HT': the spelling the template's dropdown offers."""
    return re.sub(r'\s+panels?$', '', label, flags=re.IGNORECASE) or label


def _categories():
    """``(alias table, readable list)`` for the ACTIVE nav categories."""
    table, names = {}, []
    for c in NavCategory.objects.filter(active=True):
        names.append(_short(c.label))
        for alias in (c.code, f'{c.code} panel', f'{c.code} panels', c.label, c.slug,
                      _short(c.label)):
            table.setdefault(_norm(alias), c.code)
    return table, ', '.join(names)


_ROLE = {**dict.fromkeys(('incomer', 'incomers', 'incoming', 'incomings', 'ic', 'i c'), 'incoming'),
         **dict.fromkeys(('outgoer', 'outgoers', 'outgoing', 'outgoings', 'og', 'o g'), 'outgoing'),
         **dict.fromkeys(('spare', 'spares'), 'spare')}
_ROLE_LABEL = {'incoming': 'Incomings', 'outgoing': 'Outgoings', 'spare': 'Spare'}


def _text(raw):
    if raw is None:
        return ''
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


def _convert(col, raw, categories):
    """One cell -> its value, or ValueError with a message for the operator."""
    if col.kind == 'text':
        return _text(raw)
    if col.kind in ('number', 'int'):
        try:
            value = float(raw) if isinstance(raw, (int, float)) and not isinstance(raw, bool) \
                else float(_text(raw).replace(',', ''))
        except ValueError:
            raise ValueError(f'"{_text(raw)}" is not a number') from None
        if col.kind == 'int':
            if not value.is_integer():
                raise ValueError(f'"{_text(raw)}" is not a whole number')
            return int(value)
        return value
    table, allowed = {'category': categories,
                      'role': (_ROLE, 'Incomings, Outgoings, Spare'),
                      }[col.kind]
    value = table.get(_norm(raw))
    if value is None:
        raise ValueError(f'"{_text(raw)}" is not one of {allowed}')
    return value


# --------------------------------------------------------------------------
# reading the upload
# --------------------------------------------------------------------------
class SheetError(Exception):
    """The upload cannot be read at all (wrong file type, no Meters sheet)."""


def _rows(iterator):
    """(header cells, [(sheet row number, cells), …]) — blank rows dropped."""
    header, rows = None, []
    for number, cells in enumerate(iterator, start=1):
        cells = list(cells)
        if header is None:
            header = cells
            continue
        if all(_text(c) == '' for c in cells):
            continue
        rows.append((number, cells))
        if len(rows) > MAX_ROWS:
            raise SheetError(f'More than {MAX_ROWS} rows — split the upload.')
    return header or [], rows


def read_upload(filename, data):
    """``{'Meters': (header, rows), 'Panels': (header, rows)?}`` from .xlsx or .csv."""
    name = (filename or '').lower()
    if name.endswith('.csv') or (not name.endswith(('.xlsx', '.xlsm')) and not data[:2] == b'PK'):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252')
        return {METERS: _rows(csv.reader(io.StringIO(text)))}

    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:                    # openpyxl raises a zoo of types
        raise SheetError(f'Not a readable .xlsx file ({exc}).') from None
    by_name = {_norm(ws.title): ws for ws in wb.worksheets}
    meters = by_name.get('meters')
    if meters is None:
        # a hand-made workbook: the first sheet whose header has "MFM name"
        for ws in wb.worksheets:
            if _norm(ws.title) == 'example':
                continue
            head = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if 'mfm name' in {_norm(h) for h in head}:
                meters = ws
                break
    if meters is None:
        raise SheetError('No "Meters" sheet (or any sheet with an "MFM name" column).')
    out = {METERS: _rows(meters.iter_rows(values_only=True))}
    if 'panels' in by_name:
        out[PANELS] = _rows(by_name['panels'].iter_rows(values_only=True))
    return out


# --------------------------------------------------------------------------
# the report
# --------------------------------------------------------------------------
@dataclass
class Report:
    committed: bool = False
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    changes: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)

    @property
    def ok(self):
        return not self.errors

    def error(self, sheet, row, message):
        self.errors.append({'sheet': sheet, 'row': row, 'message': message})

    def warn(self, sheet, row, message):
        self.warnings.append({'sheet': sheet, 'row': row, 'message': message})

    def change(self, entity, action, name, sheet, row, detail=''):
        counts = self.summary.setdefault(entity, {})
        counts[action] = counts.get(action, 0) + 1
        if action != 'unchanged':
            self.changes.append({'entity': entity, 'action': action, 'name': name,
                                 'sheet': sheet, 'row': row, 'detail': detail})

    def as_dict(self):
        return {'ok': self.ok, 'committed': self.committed, 'summary': self.summary,
                'errors': self.errors, 'warnings': self.warnings, 'changes': self.changes}


# --------------------------------------------------------------------------
# parse + validate (no DB)
# --------------------------------------------------------------------------
@dataclass
class MeterRow:
    row: int
    name: str
    category: str
    panel: str
    role: str | None
    table: str
    schema_code: str
    nameplate: dict

    @property
    def key(self):
        """Which meter this row means: its table when it has one, else its name."""
        return ('table', self.table) if self.table else ('name', _name_key(self.name))


@dataclass
class PanelRow:
    row: int
    name: str
    category: str | None
    sort_order: int | None
    nameplate: dict


def _parse(sheet, header, rows, columns, report, categories):
    known = {}
    for col in columns:
        known[_norm(col.label)] = col
        known.setdefault(_norm(col.key), col)
    for fname in NAMEPLATE_FIELDS:
        known.setdefault(_norm(fname), Column(fname, fname, 'number'))
    mapping = {}
    for index, head in enumerate(header):
        if _text(head) == '':
            continue
        col = known.get(_norm(head))
        if col is None:
            report.warn(sheet, 1, f'Unknown column "{_text(head)}" — ignored.')
        elif col.key in {c.key for c in mapping.values()}:
            report.error(sheet, 1, f'Column "{_text(head)}" appears twice.')
        else:
            mapping[index] = col
    mapped = {c.key for c in mapping.values()}
    for col in columns:
        if col.required and col.key not in mapped:
            report.error(sheet, 1, f'Missing column "{col.label}".')
    if not report.ok:
        return []

    parsed = []
    for number, cells in rows:
        values, nameplate, failed = {}, {}, set()
        for index, col in mapping.items():
            raw = cells[index] if index < len(cells) else None
            if _text(raw) == '':
                continue
            try:
                value = _convert(col, raw, categories)
            except ValueError as exc:
                report.error(sheet, number, f'{col.label}: {exc}.')
                failed.add(col.key)
                continue
            (nameplate if col.key in NAMEPLATE_FIELDS else values)[col.key] = value
        for col in columns:
            if col.required and col.key not in values and col.key not in failed:
                report.error(sheet, number, f'{col.label} is required.')
        parsed.append((number, values, nameplate))
    return parsed


def _meter_rows(sheets, report, categories):
    header, rows = sheets[METERS]
    return [MeterRow(row=n, name=v.get('name', ''), category=v.get('category', ''),
                     panel=v.get('panel', ''), role=v.get('role'),
                     table=v.get('table_name', ''),
                     schema_code=v.get('logger_schema', ''), nameplate=np)
            for n, v, np in _parse(METERS, header, rows, METER_COLUMNS, report, categories)]


def _panel_rows(sheets, report, categories):
    if PANELS not in sheets:
        return []
    header, rows = sheets[PANELS]
    return [PanelRow(row=n, name=v.get('name', ''), category=v.get('category'),
                     sort_order=v.get('sort_order'), nameplate=np)
            for n, v, np in _parse(PANELS, header, rows, PANEL_COLUMNS, report, categories)]


def _label(cid):
    return category_label(cid)


def _validate(meters, panels, report):
    for m in meters:
        if m.panel and not m.role:
            report.error(METERS, m.row, 'Role is required when Distribution panel is set.')
        if not m.panel and m.role:
            report.error(METERS, m.row, 'Role must be blank without a Distribution panel — an '
                                        'MFM straight under a category has no role.')

    by_meter = {}
    for m in meters:
        by_meter.setdefault(m.key, []).append(m)
    for rows in by_meter.values():
        first = rows[0]
        values, schema_seen = {}, None
        for r in rows:
            if _name_key(r.name) != _name_key(first.name):
                report.error(METERS, r.row, f'Table "{r.table}" is already the MFM '
                                            f'"{first.name}" on row {first.row}.')
            for fname, value in r.nameplate.items():
                seen = values.setdefault(fname, (value, r.row))
                if seen[0] != value:
                    report.error(METERS, r.row, f'{fname}: {value} conflicts with {seen[0]} on '
                                                f'row {seen[1]} for the same MFM.')
            if r.schema_code:
                if schema_seen is not None and r.schema_code != schema_seen[0]:
                    report.error(METERS, r.row, f'Schema "{r.schema_code}" conflicts with '
                                                f'"{schema_seen[0]}" on row {schema_seen[1]} '
                                                f'for the same MFM.')
                schema_seen = schema_seen or (r.schema_code, r.row)
        direct = [r for r in rows if not r.panel]
        placed = [r for r in rows if r.panel]
        if direct and placed:
            report.error(METERS, direct[0].row,
                         f'MFM "{first.name}" is straight under a category here but inside '
                         f'panel "{placed[0].panel}" on row {placed[0].row} — pick one.')
        if len({r.category for r in direct}) > 1:
            report.error(METERS, direct[-1].row,
                         f'MFM "{first.name}" is under two categories '
                         f'(rows {", ".join(str(r.row) for r in direct)}).')

    by_name = {}
    for key, rows in by_meter.items():
        by_name.setdefault(_name_key(rows[0].name), []).append(rows[0])
    for rows in by_name.values():
        if len(rows) > 1:
            report.warn(METERS, rows[1].row, f'Two different meters are named "{rows[0].name}" '
                                             f'(rows {rows[0].row} and {rows[1].row}).')

    by_panel = {}
    for m in meters:
        if m.panel:
            by_panel.setdefault(_name_key(m.panel), []).append(m)
    for rows in by_panel.values():
        first = rows[0]
        seen = {}
        for r in rows:
            if r.category != first.category:
                report.error(METERS, r.row, f'Panel "{r.panel}" is under {_label(first.category)} '
                                            f'on row {first.row} — a panel has one category.')
            if r.key in seen:
                report.error(METERS, r.row, f'MFM "{r.name}" is already in panel "{r.panel}" on '
                                            f'row {seen[r.key]} — one role per panel.')
            seen.setdefault(r.key, r.row)

    listed = {}
    for p in panels:
        key = _name_key(p.name)
        if key in listed:
            report.error(PANELS, p.row, f'Panel "{p.name}" is already listed on row {listed[key]}.')
            continue
        listed[key] = p.row
        used = by_panel.get(key)
        if used and p.category and p.category != used[0].category:
            report.error(PANELS, p.row, f'Nav category {_label(p.category)} does not match '
                                        f'{_label(used[0].category)} on Meters row {used[0].row}.')
        if not used and not p.category:
            report.error(PANELS, p.row, 'Nav category is required — no Meters row uses this panel.')


# --------------------------------------------------------------------------
# resolve against the DB + LoggerFast, then write
# --------------------------------------------------------------------------
@dataclass
class _Panel:
    name: str
    category: str
    sheet: str
    row: int
    sort_order: int | None = None
    nameplate: dict = field(default_factory=dict)
    group: Group | None = None


@dataclass
class _Meter:
    name: str
    row: int
    table: str
    device: dict | None
    logger_schema: LoggerSchema | None
    category: str | None            # set only for a meter straight under a category
    nameplate: dict
    mfm: MFM | None = None


def _resolve(meters, panels, report):
    by_table = {}
    if any(m.table for m in meters):
        try:
            devices = logger_devices.devices(strict=True)
        except LoggerApiError as exc:
            report.error(METERS, None, f'Cannot check table names against LoggerFast: {exc}')
            return None, None
        by_table = {d['table_name']: d for d in devices.values() if d['table_name']}
        for m in meters:
            if m.table and m.table not in by_table:
                report.error(METERS, m.row, f'Table "{m.table}" is not a LoggerFast table.')

    schema_codes = {m.schema_code for m in meters if m.schema_code}
    schemas = {schema.code: schema for schema in
               LoggerSchema.objects.filter(code__in=schema_codes, version=1)}
    for m in meters:
        if not m.schema_code:
            continue
        schema = schemas.get(m.schema_code)
        if schema is None:
            report.error(METERS, m.row, f'Logger schema "{m.schema_code}" version 1 does not exist.')
        elif not schema.active:
            report.error(METERS, m.row, f'Logger schema "{m.schema_code}" version 1 is retired.')

    plan_meters = {}
    for m in meters:
        info = plan_meters.get(m.key)
        if info is None:
            info = plan_meters[m.key] = _Meter(
                name=m.name, row=m.row, table=m.table, device=by_table.get(m.table),
                logger_schema=schemas.get(m.schema_code), category=None, nameplate={})
        elif m.schema_code and info.logger_schema is None:
            info.logger_schema = schemas.get(m.schema_code)
        if not m.panel:
            info.category = m.category
        info.nameplate.update(m.nameplate)
    for info in plan_meters.values():
        if info.device is not None:
            # the device may be linked already (possibly retired); else its table
            qs = MFM.objects.filter(logger_id=info.device['id'])
            if not qs.exists():
                qs = MFM.objects.filter(active=True, table_name=info.table)
        elif info.table:
            continue                                # unknown table, already reported
        else:
            qs = MFM.objects.filter(active=True, name__iexact=info.name)
        found = list(qs[:2])
        if len(found) > 1:
            report.error(METERS, info.row, f'Several existing meters match "{info.name}" — '
                                           f'fix them in the admin first.')
        elif found:
            info.mfm = found[0]
        if info.mfm is not None and info.category and info.mfm.memberships.filter(
                group__active=True, group__is_panel=True).exists():
            report.warn(METERS, info.row, f'MFM "{info.name}" is already in a panel and keeps '
                                          f'showing there — an upload never removes it.')

    plan_panels = {}
    for m in meters:
        if m.panel:
            plan_panels.setdefault(_name_key(m.panel),
                                   _Panel(name=m.panel, category=m.category, sheet=METERS, row=m.row))
    for p in panels:
        info = plan_panels.setdefault(_name_key(p.name), _Panel(
            name=p.name, category=p.category, sheet=PANELS, row=p.row))
        info.sort_order = p.sort_order
        info.nameplate.update(p.nameplate)
    for info in plan_panels.values():
        found = list(Group.objects.filter(active=True, name__iexact=info.name)[:2])
        if len(found) > 1:
            report.error(info.sheet, info.row, f'Several existing groups are named '
                                               f'"{info.name}" — fix them in the admin first.')
        elif found and not found[0].is_panel:
            report.error(info.sheet, info.row, f'"{info.name}" already exists as a plain group, '
                                               f'not a distribution panel.')
        elif found:
            info.group = found[0]
    return plan_meters, plan_panels


def _set(obj, fields, changes):
    """Assign ``fields`` on ``obj``; record 'field: old → new' for real changes."""
    for fname, value in fields.items():
        old = getattr(obj, fname)
        if old != value:
            setattr(obj, fname, value)
            changes.append(f'{fname}: {old if old not in (None, "") else "—"} → {value}')


def _apply_panels(plan_panels, report):
    for info in plan_panels.values():
        changes, g = [], info.group
        kind, nav = panel_axes(info.category)
        if g is None:
            g = info.group = Group.objects.create(
                name=info.name, slug=free_slug(info.name), is_panel=True,
                nav_category=nav, nav_category_source='manual', sort_order=info.sort_order or 0)
            cfg = PanelConfig.objects.create(group=g, panel_kind=kind)
            action, changes = 'create', [f'under {_label(info.category)}']
        else:
            cfg, _new = PanelConfig.objects.get_or_create(group=g)
            action = 'update'
            if resolved_category(g) != info.category:
                changes.append(f'category: {_label(resolved_category(g))} → {_label(info.category)}')
                g.nav_category, g.nav_category_source = nav, 'manual'
                cfg.panel_kind = kind
            if info.sort_order is not None:
                _set(g, {'sort_order': info.sort_order}, changes)
            g.save()
        _set(cfg, info.nameplate, changes)
        cfg.save()
        from core.equipment import ensure_panel_equipment
        ensure_panel_equipment(g)
        if action == 'update' and not changes:
            action = 'unchanged'
        report.change('panel', action, g.name, info.sheet, info.row, '; '.join(changes))


def bind_device(m, dev, changes):
    """Point meter ``m`` (unsaved changes) at LoggerFast device ``dev``: its
    table, data source and device details, pinned 'manual' so the hierarchy
    importer never rebinds it."""
    fields = {'logger_id': dev['id'], 'table_name': dev['table_name'],
              'logger_table_id': dev['table_id'], 'table_binding': 'manual',
              'protocol': dev['protocol'], 'gateway_id': dev['gateway_id'],
              'unit_id': dev['unit_id'], 'port': dev['port'],
              'logger_status': dev['status']}
    before = len(changes)
    _set(m, fields, changes)
    ds = logger_devices.data_source_for(dev['db_target_id'])
    if m.data_source_id != ds.pk:
        m.data_source = ds
        changes.append(f'data_source: → {ds.name}')
    # device details are bookkeeping; only the binding is worth reporting
    changes[before:] = [c for c in changes[before:]
                        if c.split(':')[0] in ('logger_id', 'table_name', 'data_source')]


def _apply_meters(plan_meters, report):
    for info in plan_meters.values():
        changes, m = [], info.mfm
        if m is None:
            m = info.mfm = MFM(name=info.name, slug=free_slug(info.name))
            action = 'create'
        else:
            action = 'update'
            if not m.active:
                m.active, m.retired_at, action = True, None, 'reactivate'
                if slug_owner(m.slug, m):
                    m.slug = free_slug(m.name, m)
            _set(m, {'name': info.name, 'sheet_row': info.row}, changes)
        if action == 'create':
            m.sheet_row = info.row
        if info.device is not None:
            bind_device(m, info.device, changes)
        if info.logger_schema is not None and m.logger_schema_id != info.logger_schema.pk:
            m.logger_schema = info.logger_schema
            changes.append(f'logger_schema: → {info.logger_schema.code}_v1')
        if info.category:
            _set(m, {'nav_category': info.category, 'nav_category_source': 'manual'}, changes)
        m.save()
        cfg, _new = MfmConfig.objects.get_or_create(mfm=m)
        _set(cfg, info.nameplate, changes)
        cfg.save()
        from core.equipment import ensure_mfm_equipment
        ensure_mfm_equipment(m)
        if action == 'create':
            changes.insert(0, f'under {_label(info.category)}' if info.category else 'in a panel')
        elif action == 'update' and not changes:
            action = 'unchanged'
        report.change('mfm', action, m.name, METERS, info.row, '; '.join(changes))


def _apply_memberships(meters, plan_meters, plan_panels, report):
    positions = {}
    for r in meters:
        if not r.panel:
            continue
        group = plan_panels[_name_key(r.panel)].group
        mfm = plan_meters[r.key].mfm
        slot = (group.pk, r.role)
        position = positions[slot] = positions.get(slot, -1) + 1
        label = f'{mfm.name} → {group.name} ({_ROLE_LABEL[r.role]})'
        row = Membership.objects.filter(group=group, mfm=mfm).first()
        if row is None:
            Membership.objects.create(group=group, mfm=mfm, role=r.role, position=position)
            report.change('membership', 'create', label, METERS, r.row)
        else:
            changes = []
            _set(row, {'role': r.role, 'position': position}, changes)
            if changes:
                row.save(update_fields=['role', 'position'])
            report.change('membership', 'update' if changes else 'unchanged', label,
                          METERS, r.row, '; '.join(changes))
        from core.equipment import ensure_primary_reference
        ensure_primary_reference(mfm)


def run_import(sheets, *, commit=False):
    """Parse, validate and apply an upload. Writes only when ``commit`` and the
    whole upload is clean; otherwise every write is rolled back."""
    report = Report()
    categories = _categories()
    meters = _meter_rows(sheets, report, categories)
    panels = _panel_rows(sheets, report, categories)
    # Both checking stages always run, so one upload reports EVERY error
    # (sheet rules and LoggerFast/DB lookups) instead of one stage at a time.
    _validate(meters, panels, report)
    with transaction.atomic():
        plan_meters, plan_panels = _resolve(meters, panels, report)
        if report.ok:
            _apply_panels(plan_panels, report)
            _apply_meters(plan_meters, report)
            _apply_memberships(meters, plan_meters, plan_panels, report)
        if not (commit and report.ok):
            transaction.set_rollback(True)
    report.committed = commit and report.ok
    if report.committed:
        from core.services import reset_energy_directions
        reset_energy_directions()
    return report


# --------------------------------------------------------------------------
# the template
# --------------------------------------------------------------------------
_INSTRUCTIONS = [
    ('EMS import template', 'title'),
    ('One upload builds the EMS tree. It fills the EMS nav categories listed on the Lists sheet.', ''),
    ('', ''),
    ('Meters sheet (required) — one row per meter per place it appears', 'h'),
    ('• No Distribution panel → the MFM shows directly under its Nav category. Leave Role blank.', ''),
    ('• With a Distribution panel → the panel shows under the Nav category and this MFM is its '
     'Incomings, Outgoings or Spare. The panel is created the first time its name appears.', ''),
    ('• Every row of the same panel must give the same Nav category.', ''),
    ('• A meter feeding one panel and incoming to another = two rows with the same MFM name and '
     'Table name, one per panel.', ''),
    ('• Row order = display order inside a panel\'s Incomings / Outgoings / Spares.', ''),
    ('• Table name must be a LoggerFast table; it links the device and data source.', ''),
    ('• Schema is the versioned LoggerFast field layout assigned to the MFM.', ''),
    ('', ''),
    ('Panels sheet (optional) — one row per panel', 'h'),
    ('• Only for panel nameplate values or sort order, or a panel that has no meters yet.', ''),
    ('', ''),
    ('Nameplate columns (green)', 'h'),
    ('• Optional. A blank cell leaves the stored value untouched; it never clears it.', ''),
    ('• Voltages are line-to-line in volts: 415, 11000, 33000.', ''),
    ('• Any other nameplate field can be added as an extra column headed by its field name '
     '(e.g. thd_v_limit_pct).', ''),
    ('', ''),
    ('Uploading — POST /api/import/ (form field "file"), admin login as HTTP Basic auth', 'h'),
    ('• Without commit=true it is a preview: it reports what would be created or updated and '
     'every row error, and writes nothing.', ''),
    ('• With commit=true it writes — only if there are no errors at all (all or nothing).', ''),
    ('• Re-uploading updates in place and never deletes anything. Meters are matched by Table '
     'name (or MFM name when there is none); panels by name.', ''),
    ('• Values are case-insensitive: "pcc", "PCC Panels" and "PCC" all work; "Incoming"/"IC" and '
     '"Outgoing"/"OG" are accepted for Role.', ''),
    ('• The Example sheet is ignored by the upload.', ''),
]

# Generic placeholders only — never real plant data.
_EXAMPLE_METERS = [
    ['MFM-01', 'PCC', 'Panel A', 'Incomings', 'table_mfm_01', None, 415, 1000],
    ['MFM-02', 'PCC', 'Panel A', 'Outgoings', 'table_mfm_02', None, 415],
    ['MFM-03', 'PCC', 'Panel A', 'Spare', 'table_mfm_03', None, 415],
    ['MFM-02', 'Utility', 'Panel B', 'Incomings', 'table_mfm_02', None, 415],
    ['MFM-04', 'HVAC', None, None, 'table_mfm_04', None, 415],
]
_EXAMPLE_NOTES = [
    'Panel A under PCC Panels: its incomer',
    'Panel A: an outgoer …',
    'Panel A: a spare',
    '… and the SAME meter (same table) is Panel B\'s incomer',
    'no panel: MFM-04 shows straight under HVAC Panels',
]
_EXAMPLE_PANELS = [['Panel A', 'PCC', 0, 415, 1000, None, 1600, 50, 50]]


def build_template():
    """The import template as .xlsx bytes."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    fills = {'required': PatternFill('solid', fgColor='1F4E78'),
             'optional': PatternFill('solid', fgColor='5B9BD5'),
             'nameplate': PatternFill('solid', fgColor='548235')}
    example_fill = PatternFill('solid', fgColor='FFF2CC')

    def header(ws, columns, row=1):
        for i, col in enumerate(columns, start=1):
            tone = ('required' if col.required else
                    'nameplate' if col.key in NAMEPLATE_FIELDS else 'optional')
            cell = ws.cell(row=row, column=i, value=col.label)
            cell.fill, cell.font = fills[tone], Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if col.note:
                cell.comment = Comment(col.note, 'EMS')
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(col.label) + 4)
        ws.row_dimensions[row].height = 32

    def dropdown(ws, source, column):
        dv = DataValidation(type='list', formula1=source, allow_blank=True, showErrorMessage=True,
                            errorTitle='Not in the list', error='Pick a value from the dropdown.')
        dv.add(f'{column}2:{column}{MAX_ROWS + 1}')
        ws.add_data_validation(dv)

    wb = Workbook()
    info = wb.active
    info.title = 'Instructions'
    info.column_dimensions['A'].width = 120
    for r, (text, kind) in enumerate(_INSTRUCTIONS, start=1):
        cell = info.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'title':
            cell.font = Font(bold=True, size=16)
        elif kind == 'h':
            cell.font = Font(bold=True, size=12, color='1F4E78')

    meters = wb.create_sheet(METERS)
    header(meters, METER_COLUMNS)
    meters.freeze_panes = 'B2'
    panels = wb.create_sheet(PANELS)
    header(panels, PANEL_COLUMNS)
    panels.freeze_panes = 'B2'

    example = wb.create_sheet('Example')
    example.cell(row=1, column=1, value='Ignored by the upload — shows the shape with placeholder '
                                        'names only.').font = Font(italic=True)
    header(example, METER_COLUMNS, row=2)
    for r, values in enumerate(_EXAMPLE_METERS, start=3):
        for c, value in enumerate(values, start=1):
            example.cell(row=r, column=c, value=value).fill = example_fill
        example.cell(row=r, column=len(METER_COLUMNS) + 2,
                     value=_EXAMPLE_NOTES[r - 3]).font = Font(italic=True)
    start = len(_EXAMPLE_METERS) + 5
    example.cell(row=start - 1, column=1, value='Panels sheet:').font = Font(bold=True)
    header(example, PANEL_COLUMNS, row=start)
    for r, values in enumerate(_EXAMPLE_PANELS, start=start + 1):
        for c, value in enumerate(values, start=1):
            example.cell(row=r, column=c, value=value).fill = example_fill

    lists = wb.create_sheet('Lists')
    names = [_short(c.label) for c in NavCategory.objects.filter(active=True)]
    for c, (title, values) in enumerate((
            ('Nav category', names),
            ('Role', ['Incomings', 'Outgoings', 'Spare'])), start=1):
        lists.cell(row=1, column=c, value=title).font = Font(bold=True)
        for r, value in enumerate(values, start=2):
            lists.cell(row=r, column=c, value=value)
        lists.column_dimensions[get_column_letter(c)].width = 16
    last = max(len(names), 1) + 1
    dropdown(meters, f'Lists!$A$2:$A${last}', 'B')
    dropdown(meters, 'Lists!$B$2:$B$4', 'D')
    dropdown(panels, f'Lists!$A$2:$A${last}', 'B')

    wb.active = 1                       # opens on the Meters sheet
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
