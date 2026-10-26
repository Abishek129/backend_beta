"""The spreadsheet upload behind POST /api/import/ (core.sheet_import).

Generic placeholder devices only; LoggerFast is mocked.
"""
import base64
import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from core.logger_client import LoggerApiError
from core.models import Group, LoggerSchema, Membership, MFM, MfmConfig, PanelConfig
from core.nav import build_tree
from core.sheet_import import METER_COLUMNS, PANEL_COLUMNS


def _device(n, table=True):
    return {'id': f'dev_{n}', 'name': f'DEVICE_{n}', 'protocol': 'modbus', 'gateway_id': 'gw_1',
            'unit_id': n, 'port': 502, 'status': 'connected',
            'table_name': f'table_{n:02d}' if table else '', 'table_id': f'tbl_{n}',
            'db_target_id': 'pg_main', 'ambiguous': False}


DEVICES = {d['id']: d for d in [_device(n) for n in range(1, 7)] + [_device(9, table=False)]}

M = ['MFM name', 'Nav category', 'Distribution panel', 'Role', 'Table name']


def _sheet(ws, header, rows):
    ws.append(header)
    for row in rows:
        ws.append(row)


def xlsx(meters, panels=None, header=M, panel_header=None):
    wb = Workbook()
    _sheet(wb.active, header, meters)
    wb.active.title = 'Meters'
    if panels is not None:
        _sheet(wb.create_sheet('Panels'),
               panel_header or ['Panel name', 'Nav category', 'Sort order'], panels)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BASIC = [
    ['Main Incomer', 'PCC', 'Panel A', 'Incomer', 'table_01'],
    ['Feeder 1', 'PCC', 'Panel A', 'Outgoer', 'table_02'],
    ['Feeder 2', 'PCC', 'Panel A', 'Outgoer', 'table_03'],
    ['Spare 1', 'PCC', 'Panel A', 'Spare', 'table_04'],
    ['Pump Meter', 'HVAC', None, None, 'table_05'],
]


class SheetImportCase(TestCase):
    def setUp(self):
        get_user_model().objects.create_superuser('admin', 'admin@example.com', 'pw')
        patcher = mock.patch('core.logger_devices.devices', return_value=DEVICES)
        self.devices = patcher.start()
        self.addCleanup(patcher.stop)

    def upload(self, data, commit=False, filename='ems.xlsx', auth='admin:pw'):
        headers = {}
        if auth:
            headers['HTTP_AUTHORIZATION'] = 'Basic ' + base64.b64encode(auth.encode()).decode()
        url = '/api/import/' + ('?commit=true' if commit else '')
        return self.client.post(url, {'file': SimpleUploadedFile(filename, data)}, **headers)

    def messages(self, response):
        return [f"{e['sheet']}:{e['row']} {e['message']}" for e in response.json()['errors']]


class PreviewAndCommitTests(SheetImportCase):
    def test_a_preview_reports_and_writes_nothing(self):
        r = self.upload(xlsx(BASIC))
        self.assertEqual(r.status_code, 200, r.json())
        body = r.json()
        self.assertTrue(body['ok'])
        self.assertFalse(body['committed'])
        self.assertEqual(body['summary'], {'panel': {'create': 1}, 'mfm': {'create': 5},
                                           'membership': {'create': 4}})
        self.assertEqual((Group.objects.count(), MFM.objects.count()), (0, 0))

    def test_a_commit_builds_the_tree(self):
        r = self.upload(xlsx(BASIC), commit=True)
        self.assertEqual(r.status_code, 200, r.json())
        self.assertTrue(r.json()['committed'])

        panel = Group.objects.get(name='Panel A')
        self.assertTrue(panel.is_panel)
        self.assertIsNone(panel.logger_id)
        self.assertEqual(panel.config.panel_kind, 'pcc_panel')
        self.assertEqual(dict(Membership.objects.filter(group=panel)
                              .values_list('mfm__name', 'role')),
                         {'Main Incomer': 'incoming', 'Feeder 1': 'outgoing',
                          'Feeder 2': 'outgoing', 'Spare 1': 'spare'})
        self.assertEqual(list(Membership.objects.filter(group=panel, role='outgoing')
                              .order_by('position').values_list('mfm__name', flat=True)),
                         ['Feeder 1', 'Feeder 2'])          # row order kept

        meter = MFM.objects.get(name='Main Incomer')
        self.assertEqual((meter.logger_id, meter.table_name, meter.table_binding),
                         ('dev_1', 'table_01', 'manual'))
        self.assertEqual(meter.data_source.logger_db_target_id, 'pg_main')

        pump = MFM.objects.get(name='Pump Meter')
        self.assertEqual((pump.nav_category, pump.nav_category_source), ('hvac', 'manual'))

        tree = {c['category']: c for c in build_tree()['tree']}
        node = tree['pcc']['children'][0]
        self.assertEqual([c['slug'] for c in node['children']],
                         ['overview', 'incoming', 'outgoing', 'spare'])
        self.assertEqual([n['label'] for n in tree['hvac']['children']], ['Pump Meter'])

    def test_a_second_identical_upload_changes_nothing(self):
        self.upload(xlsx(BASIC), commit=True)
        body = self.upload(xlsx(BASIC), commit=True).json()
        self.assertEqual(body['changes'], [])
        self.assertEqual(body['summary'], {'panel': {'unchanged': 1}, 'mfm': {'unchanged': 5},
                                           'membership': {'unchanged': 4}})
        self.assertEqual(MFM.objects.count(), 5)

    def test_a_reupload_updates_in_place_and_never_deletes(self):
        self.upload(xlsx(BASIC), commit=True)
        rows = [['Main Incomer (renamed)', 'Production', 'Panel A', 'Incomer', 'table_01'],
                ['Feeder 1', 'Production', 'Panel A', 'Spare', 'table_02']]
        body = self.upload(xlsx(rows), commit=True).json()
        self.assertTrue(body['ok'], body)
        panel = Group.objects.get(name='Panel A')
        self.assertEqual((panel.nav_category, panel.config.panel_kind),
                         ('production', 'distribution_panel'))
        self.assertTrue(MFM.objects.filter(name='Main Incomer (renamed)', logger_id='dev_1').exists())
        self.assertEqual(Membership.objects.get(group=panel, mfm__logger_id='dev_2').role, 'spare')
        # rows left out of this upload are untouched
        self.assertEqual(MFM.objects.count(), 5)
        self.assertEqual(Membership.objects.filter(group=panel).count(), 4)

    def test_a_meter_can_be_one_panels_outgoer_and_anothers_incomer(self):
        rows = [['Main Incomer', 'PCC', 'Panel A', 'Incomer', 'table_01'],
                ['Tie Feeder', 'PCC', 'Panel A', 'Outgoer', 'table_02'],
                ['Tie Feeder', 'Utility', 'Panel B', 'Incomer', 'table_02']]
        self.assertEqual(self.upload(xlsx(rows), commit=True).status_code, 200)
        tie = MFM.objects.get(name='Tie Feeder')
        self.assertEqual(sorted(Membership.objects.filter(mfm=tie)
                                .values_list('group__name', 'role')),
                         [('Panel A', 'outgoing'), ('Panel B', 'incoming')])

    def test_an_existing_meter_is_matched_by_its_table(self):
        MFM.objects.create(name='Old name', slug='old-name', logger_id='dev_1',
                           table_name='table_01')
        self.upload(xlsx(BASIC[:1]), commit=True)
        self.assertEqual(list(MFM.objects.values_list('name', flat=True)), ['Main Incomer'])

    def test_a_meter_without_a_table_is_matched_by_name(self):
        rows = [['Spare Slot', 'PCC', 'Panel A', 'Spare', None]]
        self.upload(xlsx(rows), commit=True)
        self.upload(xlsx(rows), commit=True)
        self.assertEqual(MFM.objects.filter(name='Spare Slot').count(), 1)
        self.assertEqual(MFM.objects.get(name='Spare Slot').logger_id, None)


class NameplateTests(SheetImportCase):
    HEADER = M + ['Nominal voltage (V)', 'Rated kVA', 'Rated kW', 'Rated current (A)',
                  'Frequency (Hz)', 'Contracted kW', 'thd_v_limit_pct']

    def test_every_meter_nameplate_column_and_extra_field_names(self):
        rows = [['Main Incomer', 'PCC', 'Panel A', 'Incomer', 'table_01',
                 415, '1,000', 800, 1390, 50, 750, 5]]
        self.assertEqual(self.upload(xlsx(rows, header=self.HEADER), commit=True).status_code, 200)
        cfg = MfmConfig.objects.get(mfm__name='Main Incomer')
        self.assertEqual((cfg.nominal_voltage_v, cfg.rated_kva, cfg.rated_kw, cfg.rated_current_a,
                          cfg.frequency_hz, cfg.contracted_kw, cfg.thd_v_limit_pct),
                         (415.0, 1000.0, 800.0, 1390.0, 50.0, 750.0, 5.0))

    def test_device_type_is_not_part_of_the_format(self):
        r = self.upload(xlsx([BASIC[0] + ['Transformer']], header=M + ['Device type']))
        self.assertEqual(r.json()['warnings'][0]['message'],
                         'Unknown column "Device type" — ignored.')

    def test_a_blank_cell_never_clears_a_stored_value(self):
        rows = [['Main Incomer', 'PCC', 'Panel A', 'Incomer', 'table_01',
                 415, None, None, None, None, None, None]]
        self.upload(xlsx(rows, header=self.HEADER), commit=True)
        rows[0][5] = None
        self.upload(xlsx(rows, header=self.HEADER), commit=True)
        self.assertEqual(MfmConfig.objects.get(mfm__name='Main Incomer').nominal_voltage_v, 415.0)

    def test_the_panels_sheet_sets_panel_nameplate_and_order(self):
        header = ['Panel name', 'Nav category', 'Sort order', 'Nominal voltage (V)',
                  'Fault level (kA)']
        panels = [['Panel A', None, 3, 11000, 25], ['Empty Board', 'APFC', None, None, None]]
        r = self.upload(xlsx(BASIC[:1], panels, panel_header=header), commit=True)
        self.assertEqual(r.status_code, 200, r.json())
        a = Group.objects.get(name='Panel A')
        self.assertEqual((a.sort_order, a.config.nominal_voltage_v, a.config.fault_level_ka),
                         (3, 11000.0, 25.0))
        empty = Group.objects.get(name='Empty Board')
        self.assertEqual((empty.is_panel, empty.nav_category), (True, 'apfc'))


class LoggerSchemaImportTests(SheetImportCase):
    def test_sheet_assigns_an_existing_active_schema(self):
        schema = LoggerSchema.objects.create(
            code='sch_EMMS7_96', version=1, name='EMMS7 96')
        header = M + ['Schema']
        row = BASIC[0] + [schema.code]
        response = self.upload(xlsx([row], header=header), commit=True)
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(MFM.objects.get(table_name='table_01').logger_schema, schema)

    def test_sheet_rejects_unknown_or_retired_schema(self):
        header = M + ['Schema']
        response = self.upload(xlsx([BASIC[0] + ['missing']], header=header), commit=True)
        self.assertEqual(response.status_code, 400)
        self.assertIn('does not exist', self.messages(response)[0])

        LoggerSchema.objects.create(
            code='retired', version=1, name='Retired', active=False)
        response = self.upload(xlsx([BASIC[0] + ['retired']], header=header), commit=True)
        self.assertEqual(response.status_code, 400)
        self.assertIn('is retired', self.messages(response)[0])


class ValidationTests(SheetImportCase):
    def assertRejected(self, meters, *expected, panels=None, header=M):
        r = self.upload(xlsx(meters, panels, header=header), commit=True)
        self.assertEqual(r.status_code, 400, r.json())
        self.assertFalse(r.json()['committed'])
        text = '\n'.join(self.messages(r))
        for fragment in expected:
            self.assertIn(fragment, text)
        self.assertEqual((Group.objects.count(), MFM.objects.count()), (0, 0))
        return text

    def test_one_bad_row_rejects_the_whole_upload(self):
        self.assertRejected(BASIC + [['Bad', 'Kitchen', None, None, None]],
                            'Meters:7 Nav category: "Kitchen" is not one of HT, AMF, PCC')

    def test_role_rules(self):
        self.assertRejected([['A', 'PCC', 'Panel A', None, 'table_01'],
                             ['B', 'PCC', None, 'Incomer', 'table_02']],
                            'Meters:2 Role is required when Distribution panel is set.',
                            'Meters:3 Role must be blank without a Distribution panel')

    def test_a_panel_has_one_category(self):
        self.assertRejected([['A', 'PCC', 'Panel A', 'Incomer', 'table_01'],
                             ['B', 'HVAC', 'Panel A', 'Outgoer', 'table_02']],
                            'Meters:3 Panel "Panel A" is under PCC Panels on row 2')

    def test_one_role_per_panel(self):
        self.assertRejected([['A', 'PCC', 'Panel A', 'Incomer', 'table_01'],
                             ['A', 'PCC', 'Panel A', 'Outgoer', 'table_01']],
                            'Meters:3 MFM "A" is already in panel "Panel A" on row 2')

    def test_straight_under_a_category_or_in_a_panel_not_both(self):
        self.assertRejected([['A', 'PCC', 'Panel A', 'Incomer', 'table_01'],
                             ['A', 'PCC', None, None, 'table_01']],
                            'Meters:3 MFM "A" is straight under a category here but inside')

    def test_a_table_belongs_to_one_meter_name(self):
        self.assertRejected([['A', 'PCC', 'Panel A', 'Incomer', 'table_01'],
                             ['B', 'PCC', 'Panel A', 'Outgoer', 'table_01']],
                            'Meters:3 Table "table_01" is already the MFM "A" on row 2')

    def test_sheet_errors_and_lookup_errors_come_back_together(self):
        self.assertRejected([['A', 'PCC', 'Panel A', None, 'no_such_table']],
                            'Meters:2 Role is required when Distribution panel is set.',
                            'Meters:2 Table "no_such_table" is not a LoggerFast table.')

    def test_unknown_table(self):
        self.assertRejected([['A', 'PCC', None, None, 'no_such_table']],
                            'Meters:2 Table "no_such_table" is not a LoggerFast table.')

    def test_bad_numbers_and_missing_columns(self):
        self.assertRejected([['A', 'PCC', None, None, None, 'lots']],
                            'Meters:2 Nominal voltage (V): "lots" is not a number.',
                            header=M + ['Nominal voltage (V)'])
        r = self.upload(xlsx([['A']], header=['MFM name']))
        self.assertIn('Meters:1 Missing column "Nav category".', self.messages(r))

    def test_a_plain_group_is_not_turned_into_a_panel(self):
        Group.objects.create(name='Panel A', slug='panel-a', logger_id='grp_x')
        r = self.upload(xlsx(BASIC[:1]), commit=True)
        self.assertEqual(r.status_code, 400)
        self.assertIn('"Panel A" already exists as a plain group', '\n'.join(self.messages(r)))

    def test_an_empty_panel_needs_a_category(self):
        r = self.upload(xlsx(BASIC[:1], [['Lonely Board', None, None]]))
        self.assertIn('Panels:2 Nav category is required — no Meters row uses this panel.',
                      self.messages(r))

    def test_loggerfast_unreachable_is_an_error_not_a_guess(self):
        self.devices.side_effect = LoggerApiError('connection refused')
        r = self.upload(xlsx(BASIC), commit=True)
        self.assertEqual(r.status_code, 400)
        self.assertIn('Cannot check table names against LoggerFast', '\n'.join(self.messages(r)))


class FormatTests(SheetImportCase):
    def test_the_users_own_headers_and_spellings_work(self):
        header = ['MFM name', 'Nav category', 'DIstribution panel', 'Role', 'table name']
        rows = [['Main Incomer', 'PCC Panels', 'Panel A', 'incoming', 'table_01'],
                ['Feeder 1', 'pcc', 'Panel A', 'OG', 'table_02']]
        r = self.upload(xlsx(rows, header=header), commit=True)
        self.assertEqual(r.status_code, 200, r.json())
        self.assertEqual(Membership.objects.count(), 2)

    def test_csv_upload(self):
        data = ('MFM name,Nav category,Distribution panel,Role,Table name\n'
                'Main Incomer,PCC,Panel A,Incomer,table_01\n'
                ',,,,\n'
                'Pump Meter,HVAC,,,table_05\n').encode()
        r = self.upload(data, commit=True, filename='ems.csv')
        self.assertEqual(r.status_code, 200, r.json())
        self.assertEqual(MFM.objects.count(), 2)

    def test_an_unknown_column_is_a_warning(self):
        r = self.upload(xlsx([BASIC[0] + ['x']], header=M + ['Remarks']))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()['warnings'][0]['message'], 'Unknown column "Remarks" — ignored.')

    def test_the_template_is_the_format(self):
        r = self.client.get('/api/import/template/')
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Instructions', 'Meters', 'Panels', 'Example', 'Lists'])
        self.assertEqual([c.value for c in wb['Meters'][1]], [c.label for c in METER_COLUMNS])
        self.assertEqual([c.value for c in wb['Panels'][1]], [c.label for c in PANEL_COLUMNS])
        # the untouched template uploads cleanly and creates nothing (Example is ignored)
        body = self.upload(r.content, commit=True).json()
        self.assertTrue(body['ok'], body)
        self.assertEqual((body['summary'], MFM.objects.count()), ({}, 0))


class AuthTests(SheetImportCase):
    def test_admin_credentials_are_required(self):
        get_user_model().objects.create_user('viewer', password='pw')
        for auth in (None, 'admin:wrong', 'viewer:pw'):
            with self.subTest(auth=auth):
                r = self.upload(xlsx(BASIC), commit=True, auth=auth)
                self.assertEqual(r.status_code, 401)
                self.assertIn('Basic', r['WWW-Authenticate'])
        self.assertEqual(MFM.objects.count(), 0)

    def test_an_admin_session_alone_is_not_enough(self):
        """CSRF-exempt, so a logged-in browser session must not authorise it."""
        self.client.login(username='admin', password='pw')
        self.assertEqual(self.upload(xlsx(BASIC), auth=None).status_code, 401)

    def test_get_is_not_allowed_and_a_file_is_required(self):
        self.assertEqual(self.client.get('/api/import/').status_code, 405)
        r = self.client.post('/api/import/', {}, HTTP_AUTHORIZATION='Basic ' +
                             base64.b64encode(b'admin:pw').decode())
        self.assertEqual(r.status_code, 400)
